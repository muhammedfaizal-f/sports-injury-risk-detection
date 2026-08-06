from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, PieChart, Reference


def build_video_report_excel(report: dict) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    # ---------- Summary ----------
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Sports Injury Risk Detection — Report"
    summary["A1"].font = Font(bold=True, size=14)

    rows = [
        ("Athlete", report.get("athlete_name", "N/A")),
        ("Video ID", report.get("video_id")),
        ("Activity Type", report.get("activity_type", "N/A")),
        ("Quality Score", report.get("quality_score", "N/A")),
        ("Risk Score", report.get("risk_score", "N/A")),
        ("Risk Category", report.get("risk_category", "N/A")),
        ("Injury Type", report.get("injury_type", "N/A")),
    ]
    for i, (label, value) in enumerate(rows, start=3):
        summary[f"A{i}"] = label
        summary[f"A{i}"].font = bold
        summary[f"B{i}"] = value

    # ---------- Joint Angles + bar chart ----------
    angles_sheet = wb.create_sheet("Joint Angles")
    angles_sheet.append(["Joint", "Value (deg)", "Max Reference (deg)"])
    for cell in angles_sheet[1]:
        cell.font = bold
    for j in report.get("joint_angles", []):
        angles_sheet.append([j["joint"], j["value"], j["max"]])

    if report.get("joint_angles"):
        n = len(report["joint_angles"])
        bar = BarChart()
        bar.title = "Joint Angles vs Reference Max"
        bar.y_axis.title = "Degrees"
        bar.x_axis.title = "Joint"
        data = Reference(angles_sheet, min_col=2, max_col=3, min_row=1, max_row=n + 1)
        cats = Reference(angles_sheet, min_col=1, min_row=2, max_row=n + 1)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.width = 18
        bar.height = 9
        angles_sheet.add_chart(bar, "E2")

    # ---------- Risk Factors + pie chart ----------
    if report.get("risk_factors"):
        risk_sheet = wb.create_sheet("Risk Factors")
        risk_sheet.append(["Factor", "Contribution"])
        for cell in risk_sheet[1]:
            cell.font = bold
        for f in report["risk_factors"]:
            risk_sheet.append([f["label"], f["contribution"]])

        n = len(report["risk_factors"])
        pie = PieChart()
        pie.title = "Risk Score Contribution by Factor"
        data = Reference(risk_sheet, min_col=2, min_row=1, max_row=n + 1)
        cats = Reference(risk_sheet, min_col=1, min_row=2, max_row=n + 1)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(cats)
        pie.width = 14
        pie.height = 9
        risk_sheet.add_chart(pie, "D2")

    # ---------- Recommendations ----------
    rec_sheet = wb.create_sheet("Recommendations")
    rec_sheet.append(["Recommendation"])
    rec_sheet["A1"].font = bold
    for r in report.get("recommendations", []):
        rec_sheet.append([r])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            try:
                sheet.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)
            except Exception:
                pass

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()