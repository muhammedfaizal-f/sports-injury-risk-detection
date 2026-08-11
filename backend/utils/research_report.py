from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, PieChart, Reference
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


def _risk_distribution_pie(risk_distribution: dict) -> BytesIO:
    fig, ax = plt.subplots(figsize=(4.5, 4.5))
    labels = [k.capitalize() for k in risk_distribution.keys()]
    values = [max(v, 0.01) for v in risk_distribution.values()]
    colors = ["#22c55e", "#f59e0b", "#f97316", "#ef4444"]
    ax.pie(values, labels=labels, autopct="%1.0f%%", colors=colors[: len(labels)], textprops={"fontsize": 8})
    ax.set_title("Risk Category Distribution", fontsize=10)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def _joint_deviation_bar(joint_deviation_frequency: dict) -> BytesIO:
    fig, ax = plt.subplots(figsize=(6, 3))
    if not joint_deviation_frequency:
        ax.text(0.5, 0.5, "No deviation data yet", ha="center", va="center")
    else:
        joints = list(joint_deviation_frequency.keys())
        counts = list(joint_deviation_frequency.values())
        ax.bar(joints, counts, color="#8b5cf6")
        ax.set_ylabel("Videos flagged")
        plt.xticks(rotation=25, ha="right", fontsize=8)
    ax.set_title("Biomechanical Deviation Frequency", fontsize=11)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def build_research_report_pdf(data: dict) -> bytes:
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 2 * cm

    def write_line(text, size=11, gap=0.6 * cm, bold=False):
        nonlocal y
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(2 * cm, y, text)
        y -= gap

    write_line("Sports Injury Risk Detection — Research Report", size=16, bold=True, gap=1 * cm)
    write_line(f"Total Athletes: {data['total_athletes']}")
    write_line(f"Total Risk Assessments: {data['total_risk_assessments']}")
    write_line(f"Total Biomechanics Analyses: {data['total_biomechanics_analyses']}")
    write_line(" ")

    pie_buf = _risk_distribution_pie(data["risk_distribution"])
    c.drawImage(ImageReader(pie_buf), 2 * cm, y - 9 * cm, width=9 * cm, height=9 * cm)
    y -= 9.5 * cm

    if y < 10 * cm:
        c.showPage()
        y = height - 2 * cm

    bar_buf = _joint_deviation_bar(data["joint_deviation_frequency"])
    c.drawImage(ImageReader(bar_buf), 2 * cm, y - 8 * cm, width=16 * cm, height=8 * cm)
    y -= 8.5 * cm

    if y < 6 * cm:
        c.showPage()
        y = height - 2 * cm

    write_line("Injury Type Distribution", size=13, bold=True)
    for injury_type, count in data["injury_type_distribution"].items():
        write_line(f"  {injury_type}: {count}", size=10)

    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer.read()


def build_research_report_excel(data: dict) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Sports Injury Risk Detection — Research Report"
    summary["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Total Athletes", data["total_athletes"]),
        ("Total Risk Assessments", data["total_risk_assessments"]),
        ("Total Biomechanics Analyses", data["total_biomechanics_analyses"]),
    ]
    for i, (label, value) in enumerate(rows, start=3):
        summary[f"A{i}"] = label
        summary[f"A{i}"].font = bold
        summary[f"B{i}"] = value

    risk_sheet = wb.create_sheet("Risk Distribution")
    risk_sheet.append(["Category", "Count"])
    for cell in risk_sheet[1]:
        cell.font = bold
    for category, count in data["risk_distribution"].items():
        risk_sheet.append([category.capitalize(), count])
    n = len(data["risk_distribution"])
    pie = PieChart()
    pie.title = "Risk Category Distribution"
    pie.add_data(Reference(risk_sheet, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
    pie.set_categories(Reference(risk_sheet, min_col=1, min_row=2, max_row=n + 1))
    risk_sheet.add_chart(pie, "D2")

    joint_sheet = wb.create_sheet("Biomechanical Deviations")
    joint_sheet.append(["Joint", "Videos Flagged"])
    for cell in joint_sheet[1]:
        cell.font = bold
    for joint, count in data["joint_deviation_frequency"].items():
        joint_sheet.append([joint, count])
    if data["joint_deviation_frequency"]:
        n = len(data["joint_deviation_frequency"])
        bar = BarChart()
        bar.title = "Biomechanical Deviation Frequency"
        bar.add_data(Reference(joint_sheet, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
        bar.set_categories(Reference(joint_sheet, min_col=1, min_row=2, max_row=n + 1))
        joint_sheet.add_chart(bar, "D2")

    injury_sheet = wb.create_sheet("Injury Type Distribution")
    injury_sheet.append(["Injury Type", "Count"])
    for cell in injury_sheet[1]:
        cell.font = bold
    for injury_type, count in data["injury_type_distribution"].items():
        injury_sheet.append([injury_type, count])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()