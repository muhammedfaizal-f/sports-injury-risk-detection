from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font


def build_video_report_excel(report: dict) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    bold = Font(bold=True)

    summary["A1"] = "Sports Injury Risk Detection — Report"
    summary["A1"].font = Font(bold=True, size=14)

    rows = [
        ("Athlete", report.get("athlete_name", "N/A")),
        ("Video ID", report.get("video_id")),
        ("Activity Type", report.get("activity_type", "N/A")),
        ("Quality Score", report.get("quality_score", "N/A")),
        ("Risk Score", report.get("risk_score", "N/A")),
        ("Risk Category", report.get("risk_category", "N/A")),
        ("Injury Type", report.get("injury_type", "N/A")),
    ]
    for i, (label, value) in enumerate(rows, start=3):
        summary[f"A{i}"] = label
        summary[f"A{i}"].font = bold
        summary[f"B{i}"] = value

    angles_sheet = wb.create_sheet("Joint Angles")
    angles_sheet.append(["Joint", "Value (deg)", "Max Reference (deg)"])
    for cell in angles_sheet[1]:
        cell.font = bold
    for j in report.get("joint_angles", []):
        angles_sheet.append([j["joint"], j["value"], j["max"]])

    rec_sheet = wb.create_sheet("Recommendations")
    rec_sheet.append(["Recommendation"])
    rec_sheet["A1"].font = bold
    for r in report.get("recommendations", []):
        rec_sheet.append([r])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            sheet.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()